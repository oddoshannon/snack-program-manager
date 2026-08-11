#!/usr/bin/env python3
"""Create a read-only matching preview for historical SNACK form exports."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_PATH = Path(__file__).resolve()
APP_ROOT = SCRIPT_PATH.parents[2]
WORKSPACE_ROOT = SCRIPT_PATH.parents[3]
DEFAULT_FORMS = WORKSPACE_ROOT / "Forms"
DEFAULT_CLIENTS = WORKSPACE_ROOT / "Clients_2026_06_01.csv"
DEFAULT_PRIVATE_OUTPUT = APP_ROOT / "backend" / ".data" / "historical-form-migration-preview.json"
DEFAULT_REVIEW_OUTPUT = APP_ROOT / "backend" / ".data" / "historical-form-match-review.csv"
DEFAULT_DUPLICATE_OUTPUT = APP_ROOT / "backend" / ".data" / "historical-form-duplicate-review.csv"
DEFAULT_SUMMARY_OUTPUT = APP_ROOT / "docs" / "HISTORICAL-FORM-MIGRATION-PREVIEW-2026-07-31.md"
REVIEW_CSV_FIELDS = (
    "sourceFile",
    "sheet",
    "rowNumber",
    "instrument",
    "programPoint",
    "sourceVersion",
    "timestamp",
    "participantName",
    "birthdate",
    "historicalIdentifier",
    "matchStatus",
    "matchMethod",
    "candidateRecordIds",
    "candidateNames",
    "candidateBirthdates",
    "possibleDuplicateExport",
    "reviewDecision",
    "reviewerNotes",
)


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def normalized(value: object) -> str:
    value_text = unicodedata.normalize("NFKD", text(value))
    value_text = "".join(character for character in value_text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", value_text.lower()).strip()


def normalized_name(value: object) -> str:
    return normalized(value)


def parse_date(value: object) -> str:
    value_text = text(value)
    if not value_text:
        return ""
    for pattern in (
        "%Y-%m-%d",
        "%Y-%m-%dT%H:%M:%S",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M:%S %p",
        "%m/%d/%y %H:%M:%S",
        "%m/%d/%y %H:%M:%S %p",
    ):
        try:
            return datetime.strptime(value_text[: len(datetime.now().strftime(pattern))], pattern).date().isoformat()
        except ValueError:
            continue
    iso_match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", value_text)
    if iso_match:
        return "-".join(iso_match.groups())
    slash_match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})", value_text)
    if slash_match:
        month, day_value, year = slash_match.groups()
        if len(year) == 2:
            year = f"20{year}"
        return f"{year}-{int(month):02d}-{int(day_value):02d}"
    return ""


def read_csv_rows(raw: bytes) -> tuple[list[str], list[dict[str, str]]]:
    decoded = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            decoded = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        decoded = raw.decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(decoded))
    headers = [text(header) for header in (reader.fieldnames or [])]
    rows = [{text(key): text(value) for key, value in row.items() if key is not None} for row in reader]
    return headers, rows


def read_form_file(path: Path) -> list[dict[str, object]]:
    sources = []
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if member.lower().endswith(".csv"):
                    headers, rows = read_csv_rows(archive.read(member))
                    sources.append({"sheet": member, "headers": headers, "rows": rows})
    elif path.suffix.lower() == ".csv":
        headers, rows = read_csv_rows(path.read_bytes())
        sources.append({"sheet": path.name, "headers": headers, "rows": rows})
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            headers = [text(value) or f"Column {index + 1}" for index, value in enumerate(values[0] if values else [])]
            rows = []
            for values_row in values[1:]:
                row = {headers[index]: text(value) for index, value in enumerate(values_row) if index < len(headers)}
                if any(row.values()):
                    rows.append(row)
            sources.append({"sheet": worksheet.title, "headers": headers, "rows": rows})
    return sources


def first_value(row: dict[str, str], aliases: tuple[str, ...], contains: tuple[str, ...] = ()) -> str:
    values = {normalized(key): text(value) for key, value in row.items()}
    for alias in aliases:
        if values.get(normalized(alias)):
            return values[normalized(alias)]
    if contains:
        for key, value in values.items():
            if value and all(fragment in key for fragment in contains):
                return value
    return ""


def client_identifier(first_name: str, last_name: str, enrollment_date: str) -> str:
    if not first_name or not last_name or not enrollment_date:
        return ""
    parsed = parse_date(enrollment_date)
    if not parsed:
        return ""
    year, month, day_value = parsed.split("-")
    return f"{first_name[0]}{last_name[0]}{month}{day_value}{year[-2:]}".upper()


def load_clients(path: Path) -> tuple[list[dict[str, str]], dict[str, list[dict[str, str]]], dict[str, list[dict[str, str]]]]:
    headers, rows = read_csv_rows(path.read_bytes())
    del headers
    clients = []
    names = defaultdict(list)
    identifiers = defaultdict(list)
    for row in rows:
        first_name = first_value(row, ("First Name",))
        last_name = first_value(row, ("Last Name",))
        full_name = first_value(row, ("Client Name",)) or f"{first_name} {last_name}".strip()
        record = {
            "recordId": first_value(row, ("Record Id", "Record ID")),
            "firstName": first_name,
            "lastName": last_name,
            "fullName": full_name,
            "dateOfBirth": parse_date(first_value(row, ("Date of Birth", "DOB"))),
            "firstAppointmentDate": parse_date(first_value(row, ("First Appt Date", "First Appointment Date"))),
        }
        record["historicalIdentifier"] = client_identifier(first_name, last_name, record["firstAppointmentDate"])
        clients.append(record)
        if normalized_name(full_name):
            names[normalized_name(full_name)].append(record)
        if record["historicalIdentifier"]:
            identifiers[record["historicalIdentifier"]].append(record)
    return clients, names, identifiers


def classify_source(file_name: str) -> tuple[str, str, str]:
    lower = file_name.lower()
    if "program enrollment" in lower:
        instrument = "Program Enrollment"
        point = "Enrollment"
    elif "parent feedback" in lower:
        instrument = "Caregiver Feedback"
        point = "Graduation"
    elif "client feedback" in lower:
        instrument = "Child Feedback"
        point = "Graduation"
    elif "initial health" in lower:
        instrument = "Health Questionnaire"
        point = "Enrollment"
    elif "final health" in lower or "post questionnaire" in lower:
        instrument = "Health Questionnaire"
        point = "Graduation"
    elif "appt" in lower:
        instrument = "Appointment Follow-Up"
        point = "Appointment"
    else:
        instrument = "Unclassified Form"
        point = "Unknown"

    if "2026" in lower:
        version = "2026"
    elif "2025" in lower:
        version = "2025"
    elif "old" in lower:
        version = "Legacy"
    else:
        version = "Undated"
    return instrument, point, version


def row_identity(row: dict[str, str]) -> dict[str, str]:
    identifier = first_value(
        row,
        (),
        contains=("office use only", "client initials", "enrollment date"),
    )
    participant_name = first_value(row, ("Participant name", "Participant Name", "Client Name", "Child Name"))
    birthdate = parse_date(first_value(row, ("Participant birthdate", "Participant Birth Date", "Date of Birth", "DOB")))
    timestamp = first_value(row, ("Timestamp", "Today's Date", "Date"))
    identifier_match = re.search(r"\b([A-Za-z]{2}\d{6})\b", re.sub(r"[^A-Za-z0-9]", "", identifier).upper())
    return {
        "identifier": identifier_match.group(1) if identifier_match else text(identifier).upper(),
        "participantName": participant_name,
        "birthdate": birthdate,
        "timestamp": timestamp,
    }


def match_identity(identity: dict[str, str], names: dict[str, list[dict[str, str]]], identifiers: dict[str, list[dict[str, str]]]) -> dict[str, object]:
    identifier = identity["identifier"]
    participant_name = identity["participantName"]
    birthdate = identity["birthdate"]
    candidates = []
    method = ""

    if identifier:
        candidates = identifiers.get(identifier, [])
        method = "initials + enrollment date"
    if not candidates and participant_name:
        candidates = names.get(normalized_name(participant_name), [])
        method = "participant name"
        if birthdate:
            matching_birthdate = [client for client in candidates if client["dateOfBirth"] == birthdate]
            if matching_birthdate:
                candidates = matching_birthdate
                method = "participant name + birthdate"

    if not identifier and not participant_name:
        return {"status": "anonymous", "method": "no client identifier", "candidates": []}
    if not candidates:
        return {"status": "unmatched", "method": method or "identifier", "candidates": []}
    if len(candidates) > 1:
        return {"status": "ambiguous", "method": method, "candidates": candidates}
    status = "exact" if method in {"participant name + birthdate", "initials + enrollment date"} else "strong"
    return {"status": status, "method": method, "candidates": candidates}


def question_count(headers: list[str]) -> int:
    metadata_terms = (
        "timestamp",
        "participant name",
        "participant birth",
        "parent guardian name",
        "parent guardian email",
        "preferred language",
        "office use only",
    )
    return sum(1 for header in headers if not any(term in normalized(header) for term in metadata_terms))


def fingerprint(row: dict[str, str], identity: dict[str, str]) -> str:
    has_stable_identity = bool(identity["identifier"] or (identity["participantName"] and identity["birthdate"]))
    serialized = json.dumps({
        normalized(key): normalized(value)
        for key, value in row.items()
        if not has_stable_identity or normalized(key) != "timestamp"
    }, sort_keys=True)
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def markdown_table_row(values: list[object]) -> str:
    return "| " + " | ".join(text(value).replace("|", "\\|").replace("\n", " ") for value in values) + " |"


def build_preview(forms_dir: Path, clients_path: Path) -> dict[str, object]:
    clients, names, identifiers = load_clients(clients_path)
    files = []
    private_rows = []
    all_fingerprints = Counter()

    for path in sorted(forms_dir.iterdir(), key=lambda item: item.name.lower()):
        if not path.is_file() or path.suffix.lower() not in {".zip", ".csv", ".xlsx"}:
            continue
        instrument, point, version = classify_source(path.name)
        for source in read_form_file(path):
            counts = Counter()
            row_records = []
            for index, row in enumerate(source["rows"], start=2):
                identity = row_identity(row)
                match = match_identity(identity, names, identifiers)
                counts[match["status"]] += 1
                row_fingerprint = fingerprint(row, identity)
                all_fingerprints[row_fingerprint] += 1
                candidates = match["candidates"]
                row_records.append({
                    "sourceFile": path.name,
                    "sheet": source["sheet"],
                    "rowNumber": index,
                    "instrument": instrument,
                    "programPoint": point,
                    "sourceVersion": version,
                    "timestamp": identity["timestamp"],
                    "participantName": identity["participantName"],
                    "birthdate": identity["birthdate"],
                    "historicalIdentifier": identity["identifier"],
                    "matchStatus": match["status"],
                    "matchMethod": match["method"],
                    "candidateClients": [
                        {
                            "recordId": candidate["recordId"],
                            "fullName": candidate["fullName"],
                            "dateOfBirth": candidate["dateOfBirth"],
                        }
                        for candidate in candidates
                    ],
                    "responseFingerprint": row_fingerprint,
                })
            private_rows.extend(row_records)
            files.append({
                "file": path.name,
                "sheet": source["sheet"],
                "instrument": instrument,
                "programPoint": point,
                "sourceVersion": version,
                "rows": len(source["rows"]),
                "columns": len(source["headers"]),
                "questionColumns": question_count(source["headers"]),
                "matchCounts": dict(counts),
                "identifierColumns": [
                    header for header in source["headers"]
                    if "participant name" in normalized(header)
                    or "participant birth" in normalized(header)
                    or "office use only" in normalized(header)
                ],
            })

    duplicate_fingerprints = {key for key, count in all_fingerprints.items() if count > 1}
    for row in private_rows:
        row["possibleDuplicateExport"] = row["responseFingerprint"] in duplicate_fingerprints
    totals = Counter(row["matchStatus"] for row in private_rows)
    return {
        "generatedAt": datetime.now().astimezone().isoformat(),
        "mode": "read-only preview; no application records changed",
        "clientSource": str(clients_path),
        "clientCount": len(clients),
        "formSource": str(forms_dir),
        "sourceCount": len(files),
        "responseCount": len(private_rows),
        "matchTotals": dict(totals),
        "possibleDuplicateResponseCount": sum(1 for row in private_rows if row["possibleDuplicateExport"]),
        "files": files,
        "rows": private_rows,
    }


def write_summary(preview: dict[str, object], output: Path) -> None:
    totals = Counter(preview["matchTotals"])
    files = preview["files"]
    lines = [
        "# Historical Form Migration Preview",
        "",
        f"Generated: {preview['generatedAt']}",
        "",
        "This is a read-only preview. No form response was imported, edited, or attached to a client profile.",
        "",
        "## Attention Needed",
        "",
        f"- **{totals['unmatched']} unmatched responses** have an identifier or participant name that did not resolve to the June 2026 client export.",
        f"- **{totals['ambiguous']} ambiguous responses** resolve to more than one client and require manual review.",
        f"- **{totals['anonymous']} anonymous responses** contain no client identifier and should remain aggregate-only unless another source can identify them.",
        f"- **{preview['possibleDuplicateResponseCount']} rows may be duplicate exports** because an identical response appears in more than one downloaded source. They must be deduplicated before import.",
        "- The two `.xlsx` questionnaire files and several zipped CSV exports overlap in questionnaire era and should not both be imported without the duplicate check.",
        "- Client Feedback, Caregiver Feedback, Dayton, and Main Appointment exports are mostly anonymous by design; they should not be guessed onto individual profiles.",
        "",
        "## Match Summary",
        "",
        markdown_table_row(["Result", "Responses", "Meaning"]),
        markdown_table_row(["---", "---:", "---"]),
        markdown_table_row(["Exact", totals["exact"], "Matched by participant name + birthdate or initials + enrollment date"]),
        markdown_table_row(["Strong", totals["strong"], "Unique participant-name match without a confirming birthdate"]),
        markdown_table_row(["Ambiguous", totals["ambiguous"], "More than one client candidate"]),
        markdown_table_row(["Unmatched", totals["unmatched"], "Identifier present but no client candidate"]),
        markdown_table_row(["Anonymous", totals["anonymous"], "No client identifier in the export"]),
        "",
        f"Source files/sheets: **{preview['sourceCount']}**  ",
        f"Historical response rows: **{preview['responseCount']}**  ",
        f"Clients available for matching: **{preview['clientCount']}**",
        "",
        "## Source Inventory",
        "",
        markdown_table_row(["Source", "Instrument", "Point", "Version", "Rows", "Questions", "Exact", "Strong", "Anonymous", "Unmatched", "Ambiguous"]),
        markdown_table_row(["---", "---", "---", "---", "---:", "---:", "---:", "---:", "---:", "---:", "---:"]),
    ]
    for file_record in files:
        counts = Counter(file_record["matchCounts"])
        source_label = file_record["file"]
        if file_record["sheet"] and file_record["sheet"] != file_record["file"]:
            source_label = f"{source_label} / {file_record['sheet']}"
        lines.append(markdown_table_row([
            source_label,
            file_record["instrument"],
            file_record["programPoint"],
            file_record["sourceVersion"],
            file_record["rows"],
            file_record["questionColumns"],
            counts["exact"],
            counts["strong"],
            counts["anonymous"],
            counts["unmatched"],
            counts["ambiguous"],
        ]))
    lines.extend([
        "",
        "## Safe Migration Order",
        "",
        "1. Manually review ambiguous and unmatched identifiers in the private preview file.",
        "2. Remove duplicate exports using the response fingerprint before creating any app records.",
        "3. Import exact matches first as locked historical versions, preserving source file, row, date, and original questions.",
        "4. Import strong name-only matches only after a person confirms them.",
        "5. Keep anonymous feedback and school/Dayton responses aggregate-only; do not assign them to a client profile.",
        "6. Re-run controlled scoring and profile-history tests after the first small import batch, before importing the remainder.",
        "",
        "## Private Detail",
        "",
        "The complete row-level candidate list is written to `backend/.data/historical-form-migration-preview.json`. A smaller decision queue is written to `backend/.data/historical-form-match-review.csv`, and possible duplicate exports are written to `backend/.data/historical-form-duplicate-review.csv`. That folder is ignored by Git because these files contain client names and identifiers.",
        "",
    ])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines), encoding="utf-8")


def candidate_values(row: dict[str, object], key: str) -> str:
    return " | ".join(text(candidate.get(key)) for candidate in row["candidateClients"] if candidate.get(key))


def review_csv_row(row: dict[str, object]) -> dict[str, object]:
    return {
        "sourceFile": row["sourceFile"],
        "sheet": row["sheet"],
        "rowNumber": row["rowNumber"],
        "instrument": row["instrument"],
        "programPoint": row["programPoint"],
        "sourceVersion": row["sourceVersion"],
        "timestamp": row["timestamp"],
        "participantName": row["participantName"],
        "birthdate": row["birthdate"],
        "historicalIdentifier": row["historicalIdentifier"],
        "matchStatus": row["matchStatus"],
        "matchMethod": row["matchMethod"],
        "candidateRecordIds": candidate_values(row, "recordId"),
        "candidateNames": candidate_values(row, "fullName"),
        "candidateBirthdates": candidate_values(row, "dateOfBirth"),
        "possibleDuplicateExport": "Yes" if row["possibleDuplicateExport"] else "No",
        "reviewDecision": "",
        "reviewerNotes": "",
    }


def write_review_csv(rows: list[dict[str, object]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=REVIEW_CSV_FIELDS)
        writer.writeheader()
        writer.writerows(review_csv_row(row) for row in rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--forms", type=Path, default=DEFAULT_FORMS)
    parser.add_argument("--clients", type=Path, default=DEFAULT_CLIENTS)
    parser.add_argument("--private-output", type=Path, default=DEFAULT_PRIVATE_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    parser.add_argument("--duplicate-output", type=Path, default=DEFAULT_DUPLICATE_OUTPUT)
    parser.add_argument("--summary-output", type=Path, default=DEFAULT_SUMMARY_OUTPUT)
    arguments = parser.parse_args()

    preview = build_preview(arguments.forms, arguments.clients)
    arguments.private_output.parent.mkdir(parents=True, exist_ok=True)
    arguments.private_output.write_text(json.dumps(preview, indent=2), encoding="utf-8")
    review_rows = [
        row for row in preview["rows"]
        if row["matchStatus"] in {"strong", "ambiguous", "unmatched"}
    ]
    duplicate_rows = [row for row in preview["rows"] if row["possibleDuplicateExport"]]
    write_review_csv(review_rows, arguments.review_output)
    write_review_csv(duplicate_rows, arguments.duplicate_output)
    write_summary(preview, arguments.summary_output)
    print(json.dumps({
        "sourceCount": preview["sourceCount"],
        "responseCount": preview["responseCount"],
        "matchTotals": preview["matchTotals"],
        "possibleDuplicateResponseCount": preview["possibleDuplicateResponseCount"],
        "reviewQueueCount": len(review_rows),
        "privateOutput": str(arguments.private_output),
        "reviewOutput": str(arguments.review_output),
        "duplicateOutput": str(arguments.duplicate_output),
        "summaryOutput": str(arguments.summary_output),
    }, indent=2))


if __name__ == "__main__":
    main()
