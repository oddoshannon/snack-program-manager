from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[2] / "outputs" / "volunteer_guide_20260811" / "Volunteering With SNACK - What to Expect.docx"

INK = "172033"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
CORAL = "D64D56"
CORAL_LIGHT = "FFF2F3"
PALE_BLUE = "E8EEF5"
MUTED = "667085"
LINE = "D7DEE8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_header_paragraph(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in (
        ("Title", 28, INK, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, BLUE_DARK, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(16)

    for list_style_name in ("List Bullet", "List Number"):
        list_style = styles[list_style_name]
        list_style.font.name = "Calibri"
        list_style.font.size = Pt(10.5)
        list_style.paragraph_format.space_after = Pt(4)
        list_style.paragraph_format.line_spacing = 1.18

    if "Small Note" not in styles:
        small = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = styles["Small Note"]
    small.font.name = "Calibri"
    small.font.size = Pt(9)
    small.font.color.rgb = RGBColor.from_string(MUTED)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.1

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "THE SNACK PROGRAM  |  VOLUNTEER GUIDE"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_p.runs[0]
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8.5)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor.from_string(BLUE_DARK)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = "The SNACK Program  •  snackprogram.org  •  (971) 202-0232"
    footer_p.runs[0].font.name = "Calibri"
    footer_p.runs[0].font.size = Pt(8.5)
    footer_p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(footer.add_paragraph())


def add_kicker(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text.upper())
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(CORAL)
    return p


def add_callout(document, heading, body, fill=CORAL_LIGHT, accent=CORAL):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "12")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    run = p.add_run(heading)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)
    run.add_break()
    body_run = p.add_run(body)
    body_run.font.color.rgb = RGBColor.from_string(INK)
    return p


def add_bullets(document, items):
    for item in items:
        # Use a literal bullet with a hanging indent. LibreOffice can otherwise
        # place Word's automatic bullet marker at the end of the preceding
        # wrapped line when it renders the DOCX to PDF.
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.26)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.18
        p.add_run("•  ")
        if isinstance(item, tuple):
            lead, rest = item
            run = p.add_run(lead)
            run.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_numbered(document, items):
    for item in items:
        p = document.add_paragraph(style="List Number")
        if isinstance(item, tuple):
            lead, rest = item
            run = p.add_run(lead)
            run.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_roles_table(document):
    rows = [
        ("Kids Cooking + Nutrition Classes", "Ingredient prep, room setup, check-in support, cleanup, and other staff-directed tasks."),
        ("Community Events", "Setup, wayfinding, activity support, hospitality, supplies, and cleanup."),
        ("Community Outreach", "Staff-supported tabling, material preparation, resource distribution, and event logistics."),
        ("Administrative Help", "Assigned office projects, organizing, data entry, and document preparation within approved access limits."),
        ("Fundraising", "Event support, thank-you projects, donor outreach support, and other approved fundraising tasks."),
        ("Marketing", "Content and design support, community promotion, and approved photography or communications projects."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(("Area", "Examples of Support")):
        cell = header.cells[index]
        set_cell_width(cell, 3240 if index == 0 else 6120)
        set_cell_shading(cell, PALE_BLUE)
        set_cell_margins(cell, 100, 120, 100, 120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    for area, support in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 3240)
        set_cell_width(cells[1], 6120)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 90, 120, 90, 120)
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.add_run(area).bold = True
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.add_run(support)
    return table


def build_document():
    document = Document()
    configure_document(document)

    # Cover
    add_kicker(document, "Volunteer orientation guide")
    title = document.add_paragraph(style="Title")
    title.add_run("Volunteering With SNACK:\nWhat to Expect")
    document.add_paragraph(
        "Application, screening, expectations, and youth safety",
        style="Subtitle",
    )
    document.add_paragraph("The SNACK Program", style="Heading 2")
    intro = document.add_paragraph(
        "Thank you for considering a volunteer role with SNACK. Volunteers help us create welcoming, organized, and accessible programs for children and families in Yamhill County. This guide explains how the process works and what we expect from every volunteer."
    )
    intro.paragraph_format.space_after = Pt(14)
    add_callout(
        document,
        "Important scope",
        "Volunteers may support classes and events only in staff-directed roles. Volunteers do not provide family nutrition education and are not placed in roles involving direct, unsupervised, or one-to-one interaction with children.",
    )
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("OUR COMMITMENT")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(CORAL)
    document.add_paragraph(
        "SNACK works to provide culturally responsive nutrition education and accessible programming. We value different identities, backgrounds, abilities, experiences, and perspectives, and we want every participant and volunteer to feel safe, respected, and included."
    )
    document.add_paragraph("Prepared August 2026", style="Small Note")

    document.add_page_break()

    # Roles
    document.add_heading("1. Ways You Can Help", level=1)
    document.add_paragraph(
        "Volunteer assignments are matched to current program needs, your interests and availability, and the level of screening and supervision appropriate for the role."
    )
    add_roles_table(document)
    document.add_heading("What volunteers do not do", level=2)
    add_bullets(document, [
        "Provide nutrition education, counseling, or individualized family services.",
        "Supervise a child alone or communicate privately with a child.",
        "Transport program participants or meet them away from an approved SNACK activity.",
        "Access participant, referral, donor, or staff information unless the role specifically requires it and access has been approved.",
        "Photograph, record, or post about participants without explicit staff authorization and the required permission.",
    ])
    add_callout(
        document,
        "Role boundaries protect everyone",
        "A volunteer can be helpful and present at a children’s program without taking on a direct service or child-supervision role. SNACK staff remain responsible for instruction, participant supervision, and program decisions.",
        fill="EEF6FB",
        accent=BLUE,
    )

    document.add_page_break()

    # Process
    document.add_heading("2. From Application to First Assignment", level=1)
    add_numbered(document, [
        ("Submit an application. ", "Tell us how you would like to help, how often you are available, whether you are interested in group volunteering, and any useful experience or questions."),
        ("Have a role-matching conversation. ", "A staff member will discuss your interests, availability, the tasks currently needed, and whether the role is a good fit."),
        ("Complete role-appropriate screening. ", "Screening may include identity or application verification, an interview, references, and—with written consent—a criminal background or sex-offender registry check. The exact steps depend on the role and its access to children, facilities, systems, money, or confidential information."),
        ("Review expectations and policies. ", "Before serving, volunteers review the Code of Conduct, Youth Protection Policy, and inclusion expectations. Separate acknowledgments or agreements may be required."),
        ("Receive approval and orientation. ", "SNACK confirms the approved role, supervision plan, access limits, and first assignment. Some roles begin with an introductory or closely supervised period."),
        ("Receive only the access you need. ", "If a role requires a SNACK email or Hub access, staff create it after approval. Access may be limited, changed, or removed when the assignment ends."),
    ])
    document.add_heading("How screening decisions are made", level=2)
    document.add_paragraph(
        "A background check is one part of screening, not the whole decision. SNACK considers the duties of the proposed role, safety needs, information gained through the application and conversation, references when used, and the requirements of current organizational policies. Screening information is handled confidentially and shared only with people who need it for the decision."
    )
    add_callout(
        document,
        "Applicants under 18",
        "Youth volunteers may have a limited set of assignments and may be asked for parent or caregiver permission. SNACK will confirm the appropriate process before placement.",
    )

    document.add_page_break()

    # Conduct
    document.add_heading("3. Volunteer Expectations", level=1)
    document.add_heading("Respect, inclusion, and professionalism", level=2)
    add_bullets(document, [
        "Treat children, caregivers, staff, partners, and other volunteers with dignity and respect.",
        "Use inclusive, culturally responsive language and remain open to feedback and learning.",
        "Follow staff direction, remain within your assigned role, and maintain appropriate personal and professional boundaries.",
        "Represent SNACK honestly and do not make commitments, promises, public statements, or purchases on SNACK’s behalf unless authorized.",
    ])
    document.add_heading("Reliability and communication", level=2)
    add_bullets(document, [
        "Arrive on time and ready for the approved assignment.",
        "Give as much notice as possible when you cannot attend; two weeks is ideal for planned absences, and same-day notice should be sent as soon as possible.",
        "Do not attend when sick. Taking time for physical or mental health is appropriate—communicate with staff so coverage can be arranged.",
        "Ask for clarification, training, or resources when you need them.",
    ])
    document.add_heading("Confidentiality and responsible access", level=2)
    add_bullets(document, [
        "Do not discuss, copy, photograph, download, or share participant, referral, donor, volunteer, or staff information outside your approved work.",
        "Use only the account and system access assigned to you; never share passwords or use another person’s account.",
        "Keep documents and screens out of public view and return or securely dispose of materials as directed.",
        "Report a mistaken disclosure, lost device, suspicious message, or other privacy concern immediately.",
    ])

    document.add_page_break()

    # Youth safety
    document.add_heading("4. Youth Safety and Boundaries", level=1)
    document.add_paragraph(
        "SNACK’s Youth Protection Policy applies whenever a volunteer’s assignment places them near children or gives them access to child-related spaces or information. Staff remain responsible for program supervision."
    )
    document.add_heading("Always", level=2)
    add_bullets(document, [
        "Stay in observable, staff-approved areas and follow the supervision plan for the activity.",
        "Use respectful, age-appropriate language and behavior.",
        "Redirect participant questions about nutrition, health, behavior, or family circumstances to SNACK staff.",
        "Follow staff instructions for check-in, photography, food safety, allergies, emergencies, and participant release.",
        "Speak up immediately if a situation feels unsafe, unclear, or outside your role.",
    ])
    document.add_heading("Never", level=2)
    add_bullets(document, [
        "Be alone with a child, arrange a private meeting, or take a child to an unobserved location.",
        "Contact a child privately by phone, text, email, social media, gaming platform, or another personal account.",
        "Transport a child or allow a child to leave with anyone except through the staff-managed release process.",
        "Use corporal punishment, humiliating language, threats, harassment, discrimination, or sexualized conduct.",
        "Take or share photographs, video, or identifying details unless staff have confirmed that the activity and permissions are approved.",
    ])
    add_callout(
        document,
        "Report concerns immediately",
        "Tell the Executive Director or designated staff member about suspected abuse, unsafe conduct, boundary violations, harassment, discrimination, or retaliation. If there is immediate danger, call 911. Do not investigate on your own or promise secrecy.",
    )

    document.add_page_break()

    # Practical guide
    document.add_heading("5. Your First Assignment", level=1)
    document.add_heading("Before you arrive", level=2)
    add_bullets(document, [
        "Confirm the date, time, location, parking or entry instructions, and staff contact.",
        "Review the assignment, dress or footwear guidance, accessibility needs, and any required training.",
        "Bring only what staff requested. Leave personal or confidential materials at home.",
        "Tell staff in advance about an accommodation or support that would help you participate successfully.",
    ])
    document.add_heading("When you arrive", level=2)
    add_bullets(document, [
        "Check in with the assigned staff coordinator before beginning work.",
        "Ask where personal items may be stored and which areas are available to volunteers.",
        "Confirm who will provide direction, what to do if you have a question, and how the assignment ends.",
        "Do not take on a new task or access new information without staff approval.",
    ])
    document.add_heading("After the assignment", level=2)
    add_bullets(document, [
        "Return keys, badges, paperwork, supplies, or devices as directed.",
        "Report an incident, near miss, privacy concern, injury, or policy question before leaving whenever possible.",
        "Share helpful feedback with the staff coordinator and confirm any future assignment separately.",
    ])
    document.add_heading("Questions and contact", level=2)
    document.add_paragraph(
        "Questions are welcome at any point. Contact Shannon Oddo, Executive Director, at director@snackprogram.org or call/text (971) 202-0232. Learn more at www.snackprogram.org."
    )
    add_callout(
        document,
        "This guide is a summary",
        "The signed Code of Conduct, Youth Protection Policy, screening documents, confidentiality rules, and role-specific instructions control if anything in this guide differs. SNACK may update assignments or expectations as programs and safety needs change.",
        fill="EEF6FB",
        accent=BLUE,
    )
    document.add_heading("Documents reflected in this guide", level=2)
    source_p = document.add_paragraph(style="Small Note")
    source_p.add_run("DEI Statement; Code of Conduct Policy & Agreement; Intern & Volunteer Screening Guidelines; and 2026 Youth Protection Policy. Reviewed August 2026.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
