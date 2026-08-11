"""Prepare the August 5 Zoho and Setmore exports for an idempotent Firestore import.

This script never connects to Firestore. It reads the approved source files and a
verified production backup, then writes a private, ignored JSON plan. Ambiguous
appointment-to-client matches are recorded but never included in the write plan.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


CLIENT_STATUS = {
    "Scheduled": "Scheduled",
    "Appts in Progress": "Active",
    "Cooking": "Active",
    "Priority Reschedule": "Needs Reschedule",
    "Parent Will Call When Ready": "Waiting on Family",
    "Too Young": "Age Limit",
    "Graduated": "Graduated",
    "On the Way Out": "Inactive",
    "Inactive": "Inactive",
}

REFERRAL_STATUS = {
    "Left Voicemail": "Left Voicemail",
    "Parent will Call Back": "Caregiver Will Call Back",
    "Not Interested": "Not Interested",
}


def normalize(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def clean(value: object) -> str:
    return str(value or "").strip()


def normalize_phone(value: object) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-10:] if len(digits) >= 10 else digits


def split_name(value: object) -> tuple[str, str]:
    parts = clean(value).split()
    if len(parts) < 2:
        return (parts[0] if parts else "", "")
    return parts[0], " ".join(parts[1:])


def read_zip_csv(path: Path) -> list[dict[str, str]]:
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name.lower().endswith(".csv")]
        if len(names) != 1:
            raise ValueError(f"Expected one CSV in {path.name}; found {len(names)}")
        with archive.open(names[0]) as stream:
            return list(csv.DictReader(line.decode("utf-8-sig") for line in stream))


def iso_date(value: object) -> str:
    raw = clean(value)
    if not raw:
        return ""
    for pattern in ("%d %b %Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Unsupported date: {raw}")


def time_details(value: object) -> tuple[str, int]:
    raw = clean(value)
    match = re.fullmatch(r"(\d{1,2}:\d{2} [AP]M)\s*-\s*(\d{1,2}:\d{2} [AP]M)", raw)
    if not match:
        raise ValueError(f"Unsupported appointment time: {raw}")
    start = datetime.strptime(match.group(1), "%I:%M %p")
    end = datetime.strptime(match.group(2), "%I:%M %p")
    minutes = int((end - start).total_seconds() // 60)
    if minutes <= 0:
        minutes += 24 * 60
    return start.strftime("%H:%M"), minutes


def clinic_type(service: object) -> str:
    value = clean(service).lower()
    if "enrollment" in value or "inscripción" in value:
        return "Enrollment"
    if "nutrition education appointment" in value or "educación nutricional" in value:
        return "Nutrition Education"
    return ""


def exact_person_candidates(row: dict, people: list[dict]) -> list[dict]:
    name_key = normalize(row.get("Customer name"))
    candidates = [person for person in people if person["nameKey"] == name_key]
    if len(candidates) <= 1:
        return candidates
    phone = normalize_phone(row.get("Phone"))
    email = clean(row.get("Email")).lower()
    dob = iso_date(row.get("Child's DOB")) if clean(row.get("Child's DOB")) else ""
    narrowed = [
        person for person in candidates
        if (phone and person["phoneKey"] == phone)
        or (email and person["emailKey"] == email)
        or (dob and person["dateOfBirth"] == dob)
    ]
    return narrowed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", required=True)
    parser.add_argument("--backup", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    source_dir = Path(args.source_dir).resolve()
    backup_path = Path(args.backup).resolve()
    output_path = Path(args.out).resolve()
    backup = json.loads(backup_path.read_text())
    collections = backup["collections"]

    zoho_clients = read_zip_csv(source_dir / "Clients_2026_08_05.zip")
    zoho_referrals = read_zip_csv(source_dir / "Referrals_2026_08_05.zip")
    existing_clients = collections.get("clients", [])
    existing_referrals = collections.get("referrals", [])
    existing_appointments = collections.get("appointments", [])
    existing_client_zoho = {row["data"].get("zohoRecordId") for row in existing_clients}
    existing_referral_zoho = {row["data"].get("zohoRecordId") for row in existing_referrals}

    clients_to_create = []
    for row_number, row in enumerate(zoho_clients, start=2):
        if row["Record Id"] in existing_client_zoho:
            continue
        first_name, last_name = split_name(row.get("Client Name"))
        clients_to_create.append({
            "id": row["Record Id"],
            "rowNumber": row_number,
            "firstName": first_name,
            "lastName": last_name,
            "parentName": clean(row.get("Parent Name")),
            "dateOfBirth": "",
            "gender": "",
            "phone": clean(row.get("Mobile")),
            "email": clean(row.get("Email")).lower(),
            "preferredLanguage": clean(row.get("Preferred Language")),
            "preferredContactMethod": "",
            "referralType": "",
            "referralSource": "",
            "referralDate": "",
            "firstContactDate": "",
            "mostRecentContactDate": clean(row.get("Most Recent Contact Date")),
            "firstAppointmentDate": "",
            "mostRecentAppointmentDate": "",
            "lastAppointmentDate": "",
            "currentLesson": "",
            "addressStreet": "",
            "addressCity": "",
            "addressState": "",
            "addressZip": "",
            "emailOptOut": False,
            "textOptOut": False,
            "ycco": False,
            "hrsn": False,
            "assessmentScore": None,
            "willingnessScore": None,
            "notes": "",
            "status": CLIENT_STATUS[row["Status"]],
            "siblingIds": [],
            "zohoRecordId": row["Record Id"],
        })

    # Build the exact client lookup from existing live records plus approved new rows.
    people = []
    for row in existing_clients:
        data = row["data"]
        people.append({
            "id": row["id"],
            "name": f"{data.get('firstName', '')} {data.get('lastName', '')}".strip(),
            "nameKey": normalize(f"{data.get('firstName', '')} {data.get('lastName', '')}"),
            "phoneKey": normalize_phone(data.get("phone")),
            "emailKey": clean(data.get("email")).lower(),
            "dateOfBirth": clean(data.get("dateOfBirth")),
            "new": False,
        })
    for data in clients_to_create:
        people.append({
            "id": data["id"],
            "name": f"{data['firstName']} {data['lastName']}".strip(),
            "nameKey": normalize(f"{data['firstName']} {data['lastName']}"),
            "phoneKey": normalize_phone(data["phone"]),
            "emailKey": data["email"],
            "dateOfBirth": "",
            "new": True,
        })

    network_providers = []
    for network in collections.get("referralNetwork", []):
        for provider in network["data"].get("providers", []):
            network_providers.append({
                "networkId": network["id"],
                "networkName": clean(network["data"].get("name")),
                "providerId": clean(provider.get("id")),
                "providerName": clean(provider.get("name")),
                "nameKey": normalize(provider.get("name")),
            })

    referrals_to_create = []
    unresolved_referral_links = []
    for row_number, row in enumerate(zoho_referrals, start=2):
        if row["Record Id"] in existing_referral_zoho:
            continue
        first_name, last_name = split_name(row.get("Referral Name"))
        provider_matches = [p for p in network_providers if p["nameKey"] == normalize(row.get("Referring Provider"))]
        provider_links = []
        if len(provider_matches) == 1:
            provider = provider_matches[0]
            provider_links = [{k: provider[k] for k in ("networkId", "networkName", "providerId", "providerName")}]
        elif clean(row.get("Referring Provider")):
            unresolved_referral_links.append({"rowNumber": row_number, "zohoRecordId": row["Record Id"], "reason": "Provider name was not an exact unique network match."})
        referrals_to_create.append({
            "id": row["Record Id"],
            "rowNumber": row_number,
            "firstName": first_name,
            "lastName": last_name,
            "parentName": clean(row.get("Parent Name")),
            "dateOfBirth": "",
            "gender": "",
            "phone": clean(row.get("Mobile")),
            "email": "",
            "preferredLanguage": clean(row.get("Preferred Language")),
            "preferredContactMethod": "",
            "referralType": "Provider Referral",
            "referralSource": clean(row.get("Referring Provider")),
            "referralDate": clean(row.get("Referral Date")),
            "firstContactDate": "",
            "mostRecentContactDate": clean(row.get("Most Recent Contact Date")),
            "firstAppointmentDate": "",
            "mostRecentAppointmentDate": "",
            "lastAppointmentDate": "",
            "currentLesson": "",
            "addressStreet": "",
            "addressCity": "",
            "addressState": "",
            "addressZip": "",
            "emailOptOut": False,
            "textOptOut": False,
            "ycco": False,
            "hrsn": False,
            "assessmentScore": None,
            "willingnessScore": None,
            "notes": "",
            "status": REFERRAL_STATUS[row["Referral Status"]],
            "siblingIds": [],
            "providerLinks": provider_links,
            "zohoRecordId": row["Record Id"],
        })

    workbook = load_workbook(source_dir / "Appointments.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    setmore_rows = [dict(zip(headers, values)) | {"rowNumber": row_number} for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2)]
    imported_row_numbers = set()
    for row in existing_appointments:
        data = row["data"]
        if data.get("importedFrom") != "Setmore appointment export":
            continue
        values = data.get("rowNumbers") if isinstance(data.get("rowNumbers"), list) else [data.get("rowNumber")]
        imported_row_numbers.update(int(value) for value in values if value)

    exact_existing = defaultdict(list)
    for row in existing_appointments:
        data = row["data"]
        for client_id in data.get("clientIds") or [data.get("clientId")]:
            if client_id:
                exact_existing[(client_id, clean(data.get("appointmentDate")), clean(data.get("appointmentTime")))].append(row["id"])

    candidates = []
    unresolved_appointments = []
    excluded_non_clinic = []
    already_reconciled = []
    for row in setmore_rows:
        row_number = row["rowNumber"]
        appt_type = clinic_type(row.get("Service/class/event"))
        if not appt_type:
            excluded_non_clinic.append({"rowNumber": row_number, "service": clean(row.get("Service/class/event")), "status": clean(row.get("Status"))})
            continue
        if row_number in imported_row_numbers:
            already_reconciled.append({"rowNumber": row_number, "reason": "Previously imported Setmore row."})
            continue
        matches = exact_person_candidates(row, people)
        if len(matches) != 1:
            unresolved_appointments.append({"rowNumber": row_number, "bookingId": clean(row.get("Booking ID")), "reason": f"Expected one exact client match; found {len(matches)}."})
            continue
        client = matches[0]
        date = iso_date(row.get("Appointment date"))
        start, duration = time_details(row.get("Appointment time"))
        if exact_existing[(client["id"], date, start)]:
            already_reconciled.append({"rowNumber": row_number, "reason": "An appointment already exists for the exact client, date, and time."})
            continue
        candidates.append({
            "rowNumber": row_number,
            "bookingId": clean(row.get("Booking ID")),
            "clientId": client["id"],
            "clientName": client["name"],
            "appointmentDate": date,
            "appointmentTime": start,
            "durationMinutes": duration,
            "appointmentType": appt_type,
            "status": "Canceled" if clean(row.get("Status")).lower().startswith("cancel") else "Scheduled",
            "lesson": "Enrollment" if appt_type == "Enrollment" else "",
            "staffMember": clean(row.get("Team member")),
            "notes": clean(row.get("Comments ")),
            "publicBookingServiceLabel": clean(row.get("Service/class/event")),
            "mergeKey": "|".join([date, start, clean(row.get("Team member")), normalize_phone(row.get("Phone")) or clean(row.get("Email")).lower(), normalize(row.get("Parent's Name"))]),
            "sibling": "sibling" in clean(row.get("Service/class/event")).lower(),
        })

    # Merge only exact sibling slots sharing a contact identity; ordinary same-time
    # appointments remain separate.
    grouped = defaultdict(list)
    for row in candidates:
        key = row["mergeKey"] if row["sibling"] else f"row:{row['rowNumber']}"
        grouped[key].append(row)

    appointments_to_create = []
    for rows in grouped.values():
        rows.sort(key=lambda item: item["rowNumber"])
        first = rows[0]
        digest = hashlib.sha256("|".join(item["bookingId"] for item in rows).encode()).hexdigest()[:24]
        appointments_to_create.append({
            "id": f"setmore-{digest}",
            "clientId": first["clientId"],
            "clientIds": [item["clientId"] for item in rows],
            "clientName": first["clientName"],
            "clientNames": [item["clientName"] for item in rows],
            "appointmentDate": first["appointmentDate"],
            "appointmentTime": first["appointmentTime"],
            "durationMinutes": first["durationMinutes"],
            "appointmentType": first["appointmentType"],
            "status": "Canceled" if all(item["status"] == "Canceled" for item in rows) else "Scheduled",
            "lesson": first["lesson"],
            "staffMember": first["staffMember"],
            "notes": "\n".join(dict.fromkeys(item["notes"] for item in rows if item["notes"])),
            "publicBookingServiceLabel": first["publicBookingServiceLabel"],
            "rowNumber": first["rowNumber"],
            "rowNumbers": [item["rowNumber"] for item in rows],
            "sourceBookingId": first["bookingId"],
            "sourceBookingIds": [item["bookingId"] for item in rows],
            "importSource": "Setmore appointment export",
            "importedFrom": "Setmore appointment export",
        })

    plan = {
        "format": "snack-approved-real-data-import-v1",
        "sourceCutoff": "2026-08-05",
        "preparedAt": datetime.now().astimezone().isoformat(),
        "backupPath": str(backup_path),
        "summary": {
            "sourceZohoClients": len(zoho_clients),
            "existingZohoClients": len(existing_client_zoho & {row['Record Id'] for row in zoho_clients}),
            "clientsToCreate": len(clients_to_create),
            "sourceZohoReferrals": len(zoho_referrals),
            "existingZohoReferrals": len(existing_referral_zoho & {row['Record Id'] for row in zoho_referrals}),
            "referralsToCreate": len(referrals_to_create),
            "sourceSetmoreRows": len(setmore_rows),
            "previouslyReconciledClinicRows": len([row for row in already_reconciled if "Previously" in row["reason"]]),
            "existingExactAppointmentRows": len([row for row in already_reconciled if "already exists" in row["reason"]]),
            "appointmentsToCreate": len(appointments_to_create),
            "appointmentSourceRowsToCreate": sum(len(row["rowNumbers"]) for row in appointments_to_create),
            "unresolvedAppointmentRows": len(unresolved_appointments),
            "excludedNonClinicRows": len(excluded_non_clinic),
            "unresolvedReferralProviderLinks": len(unresolved_referral_links),
        },
        "clientsToCreate": clients_to_create,
        "referralsToCreate": referrals_to_create,
        "appointmentsToCreate": appointments_to_create,
        "unresolvedAppointments": unresolved_appointments,
        "unresolvedReferralProviderLinks": unresolved_referral_links,
        "alreadyReconciledAppointments": already_reconciled,
        "excludedNonClinicRows": excluded_non_clinic,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(plan, indent=2) + "\n")
    output_path.chmod(0o600)
    print(json.dumps(plan["summary"], indent=2))


if __name__ == "__main__":
    main()
